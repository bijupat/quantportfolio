import logging
from pathlib import Path
from datetime import datetime
from django.core.files import File
from forecasting.models import ReportArtifact
import pandas as pd
import numpy as np

logger = logging.getLogger(__name__)


def save_portfolio_report_pdf(
    ranking_df: pd.DataFrame,
    weights: dict,
    tiers: dict,
    metrics: dict = None,
    raw_scores: dict = None,
    model_name: str = "Unknown",
    portfolio=None,
    composite_run=None,
    source: str = ReportArtifact.Source.COMPOSITE,
) -> ReportArtifact:
    """Generates a styled PDF report for the portfolio and saves it as a ReportArtifact.

    Args:
        portfolio: The portfolio.models.Portfolio this report was generated for, if any.
        composite_run: A representative forecasting.models.CompositeScore row for this
            run (see save_composite_scores_to_db) — links the artifact back to the
            scoring config/date it came from, if any.
        source: Which pipeline produced this report — one of
            ReportArtifact.Source. Defaults to COMPOSITE since run_composite.py
            is this function's original/primary caller; forecasting/management/
            commands/predict.py passes source=ReportArtifact.Source.PREDICT
            explicitly so its own single-model PDF reports are distinguishable
            in the Generated Reports table (see ReportArtifact's docstring).
    """
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.lib.units import mm
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    except ImportError:
        logger.warning("reportlab not installed.")
        return None

    import tempfile
    import os

    NAVY = colors.HexColor("#0A1628")
    DEEP_BLUE = colors.HexColor("#0D2B55")
    CYAN = colors.HexColor("#00B4D8")
    TEAL = colors.HexColor("#00897B")
    LT_GREY = colors.HexColor("#F5F7FA")
    MID_GREY = colors.HexColor("#90A4AE")
    DK_GREY = colors.HexColor("#37474F")
    GREEN = colors.HexColor("#43A047")
    AMBER = colors.HexColor("#FF8F00")
    CORAL = colors.HexColor("#EF5350")
    WHITE = colors.white

    def PS(name, **kw): return ParagraphStyle(name, **kw)
    title_st = PS("ts", fontName="Helvetica-Bold", fontSize=14, textColor=WHITE, alignment=1)
    head_st = PS("hs", fontName="Helvetica-Bold", fontSize=10, textColor=DEEP_BLUE, spaceBefore=10, spaceAfter=4)
    cell_b_st = PS("cb", fontName="Helvetica-Bold", fontSize=9, textColor=DK_GREY, leading=13)

    def page_tmpl(canv, doc):
        w, h = A4
        canv.saveState()
        canv.setFillColor(NAVY)
        canv.rect(0, h-14*mm, w, 14*mm, stroke=0, fill=1)
        canv.setFillColor(CYAN)
        canv.rect(0, h-15.5*mm, w, 1.5*mm, stroke=0, fill=1)
        canv.setFillColor(WHITE)
        canv.setFont("Helvetica-Bold", 8)
        canv.drawString(15*mm, h-9.5*mm, "QuantPortfolioAI")
        canv.drawString(90*mm, h-9.5*mm, f"Model : {model_name}")
        canv.restoreState()

    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
        doc = SimpleDocTemplate(tmp.name, pagesize=A4, leftMargin=20*mm, rightMargin=20*mm, topMargin=20*mm, bottomMargin=16*mm)
        story = []

        title_tbl = Table([[Paragraph("Predicted 30-Day Return Ranking", title_st)]], colWidths=[460])
        title_tbl.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,-1), DEEP_BLUE), ("ROUNDEDCORNERS", [5]), ("TOPPADDING", (0,0), (-1,-1), 10), ("BOTTOMPADDING", (0,0), (-1,-1), 10)]))
        story.append(title_tbl)
        story.append(Spacer(1, 12))

        story.append(Paragraph("Score Ranking", head_st))
        tier_colors = {"BUY": GREEN, "HOLD": AMBER, "AVOID": CORAL}
        ranking_rows = [[
            Paragraph("<b>Rank</b>", PS("rh", fontName="Helvetica-Bold", fontSize=8.5, textColor=WHITE)),
            Paragraph("<b>Symbol</b>", PS("rh2", fontName="Helvetica-Bold", fontSize=8.5, textColor=WHITE)),
            Paragraph("<b>Score</b>", PS("rh3", fontName="Helvetica-Bold", fontSize=8.5, textColor=WHITE, alignment=2)),
            Paragraph("<b>Tier</b>", PS("rh4", fontName="Helvetica-Bold", fontSize=8.5, textColor=WHITE, alignment=1)),
        ]]

        for rank, (sym, row) in enumerate(ranking_df.iterrows(), 1):
            tier = tiers.get(sym, "AVOID")
            score = row["norm_score"]
            tcol = tier_colors.get(tier, CORAL)
            ranking_rows.append([
                Paragraph(str(rank), PS(f"rc{rank}", fontName="Helvetica-Bold", fontSize=9, textColor=CYAN, alignment=1)),
                Paragraph(sym, cell_b_st),
                Paragraph(f"{score:.4f}", PS(f"rs{rank}", fontName="Helvetica-Bold", fontSize=9, textColor=DK_GREY, alignment=2)),
                Paragraph(tier, PS(f"rt{rank}", fontName="Helvetica-Bold", fontSize=8, textColor=WHITE, alignment=1, backColor=tcol, borderPad=3)),
            ])

        ranking_tbl = Table(ranking_rows, colWidths=[40, 200, 80, 140])
        row_bgs = [LT_GREY if i % 2 == 0 else WHITE for i in range(len(ranking_rows)-1)]
        ranking_tbl.setStyle(TableStyle([
            ("BACKGROUND", (0,0), (-1,0), DEEP_BLUE), ("LINEABOVE", (0,0), (-1,0), 1.5, CYAN),
            ("ROWBACKGROUNDS", (0,1), (-1,-1), row_bgs), ("GRID", (0,0), (-1,-1), 0.3, MID_GREY),
            ("VALIGN", (0,0), (-1,-1), "MIDDLE"), ("LEFTPADDING", (0,0), (-1,-1), 8), ("RIGHTPADDING", (0,0), (-1,-1), 8)
        ]))
        story.append(ranking_tbl)

        doc.build(story, onFirstPage=page_tmpl, onLaterPages=page_tmpl)
    artifact = ReportArtifact.objects.create(
        kind=ReportArtifact.PDF, portfolio=portfolio, composite_run=composite_run, source=source
    )
    now_str = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    with open(tmp.name, 'rb') as f:
        artifact.file.save(f"{now_str}_portfolio_report.pdf", File(f))
    os.remove(tmp.name)
    return artifact


def _is_low_quality(detail: dict) -> bool:
    """Returns True if a composite_results entry's data_quality dict flags reduced confidence.

    Mirrors the same three conditions run_composite.py checks for its console
    warning (quality_no_data / technical_no_data / partial_ensemble), so the
    persisted Excel report shows the same signal the person already saw on
    screen when the run finished — previously this information existed in
    memory (compute_composite_scores' data_quality dict) but was dropped
    entirely once results reached the Excel writer; a stock scored on
    insufficient history looked identical to a fully-scored one in the saved
    report. Tolerant of composite_results predating this field (e.g. results
    passed in from a caller that doesn't supply data_quality) — such rows are
    simply never flagged rather than raising.
    """
    dq = detail.get("data_quality") or {}
    return bool(dq.get("quality_no_data") or dq.get("technical_no_data") or dq.get("partial_ensemble"))


def save_score_breakdown_excel(
    composite_results: dict,
    portfolio_holdings: list,
    tiers: dict,
    model_names: list,
    model_weights: list,
    weights_used: dict,
    portfolio_amount: float,
    portfolio=None,
    composite_run=None,
) -> ReportArtifact:
    """
    Generates the multi-sheet Score Breakdown and Portfolio Excel workbook
    and saves it as a ReportArtifact.

    Args:
        portfolio: The portfolio.models.Portfolio this report was generated for, if any.
        composite_run: A representative forecasting.models.CompositeScore row for this
            run, if any.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        logger.warning("openpyxl not installed.")
        return None

    import tempfile
    import os

    NAVY = "FF1E3A5F"
    NAVY_LIGHT = "FF2A5080"
    WHITE = "FFFFFFFF"
    GREEN_HDR = "FF1B5E20"
    AMBER_HDR = "FFE65100"
    RED_HDR = "FFB71C1C"
    GREEN_FILL = "FFE8F5E9"
    AMBER_FILL = "FFFDF3E3"
    RED_FILL = "FFFCE4E4"
    GREEN_SCORE = "FF43A047"
    RED_SCORE = "FFEF5350"
    GREY_FONT = "FF607D8B"
    STRIPE_EVEN = "FFF5F7FA"
    STRIPE_ODD = "FFFFFFFF"
    LOW_QUALITY_FILL = "FFFFF3CD"   # pale amber — flags reduced-confidence rows
    LOW_QUALITY_FONT = "FF8A6D00"

    def _hdr_font(bold=True, color=WHITE, size=9): return Font(name="Arial", bold=bold, color=color, size=size)
    def _body_font(bold=False, color="FF0D1B2A", size=9): return Font(name="Arial", bold=bold, color=color, size=size)
    def _fill(argb): return PatternFill("solid", fgColor=argb)
    def _centre(): return Alignment(horizontal="center", vertical="center", wrap_text=True)
    def _left(): return Alignment(horizontal="left", vertical="center", wrap_text=False)

    thin = Side(style="thin", color="FFBDBDBD")
    thin_border = Border(left=thin, right=thin, top=thin, bottom=thin)

    first_d = next(iter(composite_results.values()), {})
    pm_keys = list(first_d.get("per_model_scores", {}).keys())
    pm_short = [m.replace("transformer_", "") for m in pm_keys]
    total_w = sum(model_weights) or 1
    norm_mw = [w / total_w for w in model_weights]

    sorted_results = sorted(composite_results.items(), key=lambda x: -x[1]["composite_score"])

    wb = Workbook()

    # --- Sheet 1: Score Breakdown ---
    ws = wb.active
    ws.title = "Score Breakdown"

    now_str = datetime.now().strftime("%d %b %Y  %H:%M")
    low_quality_total = sum(1 for _, d in composite_results.items() if _is_low_quality(d))
    ws.append([
        f"QuantPortfolioAI — Full Score Breakdown", f"Generated: {now_str}",
        f"Symbols: {len(composite_results)}", f"Capital: Rs. {portfolio_amount:,.0f}",
        f"Reduced-confidence: {low_quality_total}" if low_quality_total else "",
    ])
    ws.row_dimensions[1].height = 16
    ws.append([])

    HDR_ROW = 3
    # Column count here MUST stay in lockstep with the row_vals column count
    # built per-row below (both derive from the same "how many per-model
    # columns" decision) — previously this was two independently-written
    # `X if pm_keys else Y` expressions relying on Python's `+=` precedence
    # matching by coincidence; a one-line edit to either without the other
    # would silently misalign every column after "Composite" with no error
    # from openpyxl. n_transformer_cols is now computed once and reused by
    # both the header and the per-row builder so they cannot drift apart,
    # and a Data Quality column is appended (also counted here) to surface
    # the low_quality flag per row.
    n_transformer_cols = (len(pm_keys) + 1) if pm_keys else 1

    grp_row = ["Rank", "Symbol", "Composite"]
    grp_row += ["Transformer"] * n_transformer_cols
    grp_row += ["Financial", "Technical", "Tier", "Data Quality"]
    ws.append(grp_row)

    ind_row = ["", "", ""]
    ind_row += (pm_short + ["Weighted"]) if pm_keys else [""]
    ind_row += ["(Piotroski+Altman)", "(TSI)", "", ""]
    ws.append(ind_row)

    n_cols_total = len(grp_row)
    for r in (HDR_ROW, HDR_ROW + 1):
        for c in range(1, n_cols_total + 1):
            cell = ws.cell(r, c)
            cell.font = _hdr_font()
            cell.fill = _fill(NAVY)
            cell.alignment = _centre()
            cell.border = thin_border
        ws.row_dimensions[r].height = 28

    if pm_keys:
        ws.merge_cells(start_row=HDR_ROW, start_column=4, end_row=HDR_ROW, end_column=4 + len(pm_keys))
        merged_cell = ws.cell(HDR_ROW, 4)
        merged_cell.fill = _fill(NAVY_LIGHT)
        merged_cell.alignment = _centre()

    tier_row_fill = {"BUY": GREEN_FILL, "HOLD": AMBER_FILL, "AVOID": RED_FILL}
    tier_font_clr = {"BUY": GREEN_HDR, "HOLD": AMBER_HDR, "AVOID": RED_HDR}

    for rank, (sym, d) in enumerate(sorted_results, 1):
        tier = tiers.get(sym, "AVOID")
        pm = d.get("per_model_scores", {})
        low_quality = _is_low_quality(d)

        row_vals = [rank, sym, d["composite_score"]]
        row_vals += ([pm.get(m, None) for m in pm_keys] + [d["transformer_score"]]) if pm_keys else [d["transformer_norm"]]
        row_vals += [d["quality_score"], d["technical_score"], tier, "LOW" if low_quality else "OK"]

        ws.append(row_vals)
        xl_row = ws.max_row
        row_bg = _fill(tier_row_fill.get(tier, STRIPE_ODD))
        default_stripe = _fill(STRIPE_EVEN if rank % 2 == 0 else STRIPE_ODD)

        for col_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(xl_row, col_idx)
            cell.border = thin_border
            cell.fill = default_stripe

            if col_idx == 1:
                cell.font = _body_font(color=GREY_FONT)
                cell.alignment = _centre()
            elif col_idx == 2:
                cell.font = _body_font(bold=True)
                cell.alignment = _left()
            elif col_idx == 3:
                cell.number_format = "0.0000"
                cell.alignment = _centre()
                cell.font = _body_font(bold=True, color=GREEN_SCORE if val >= 0.6 else (RED_SCORE if val < 0.4 else "FF0D1B2A"))
            elif col_idx <= 3 + n_transformer_cols:
                cell.number_format = "+0.0000;-0.0000;0.0000"
                cell.alignment = _centre()
                cell.font = _body_font(color="FF1B5E20" if val > 0 else RED_SCORE) if isinstance(val, float) else _body_font()
            elif col_idx <= n_cols_total - 2:
                cell.number_format = "0.0000"
                cell.alignment = _centre()
                cell.font = _body_font()
            elif col_idx == n_cols_total - 1:
                cell.font = Font(name="Arial", bold=True, color=tier_font_clr.get(tier, GREY_FONT), size=9)
                cell.alignment = _centre()
            else:
                # Data Quality column — flagged rows get the amber fill/font
                # regardless of tier, so a BUY-tier stock with insufficient
                # history is still visibly distinct from a fully-scored BUY.
                cell.font = Font(name="Arial", bold=True, color=LOW_QUALITY_FONT if low_quality else GREEN_HDR, size=9)
                cell.alignment = _centre()

        for col_idx in range(1, n_cols_total + 1):
            if col_idx == n_cols_total:
                if low_quality:
                    ws.cell(xl_row, col_idx).fill = _fill(LOW_QUALITY_FILL)
            elif col_idx != n_cols_total - 1:
                ws.cell(xl_row, col_idx).fill = row_bg

    col_widths = [5, 22, 11] + ([13] * n_transformer_cols) + [13, 13, 8, 12]
    for col_idx, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "C5"
    ws.auto_filter.ref = f"A4:{get_column_letter(n_cols_total)}{ws.max_row}"

    # --- Sheet 2: Portfolio ---
    ws2 = wb.create_sheet("Portfolio")
    ws2.append([f"Portfolio Allocation — BUY Tier", f"Generated: {now_str}", f"Capital: Rs. {portfolio_amount:,.0f}", f"Holdings: {len(portfolio_holdings)}"])
    ws2.row_dimensions[1].height = 16
    ws2.append([])

    port_hdr = ["Symbol", "Alloc %", "Rs. Alloc", "Price (Rs.)", "Qty", "Composite", "Transformer", "Financial", "Technical"]
    ws2.append(port_hdr)
    hdr_row = ws2.max_row
    for c in range(1, len(port_hdr) + 1):
        cell = ws2.cell(hdr_row, c)
        cell.font = _hdr_font()
        cell.fill = _fill(NAVY)
        cell.alignment = _centre()
        cell.border = thin_border
    ws2.row_dimensions[hdr_row].height = 24

    total_alloc = 0.0
    for h in sorted(portfolio_holdings, key=lambda x: -x.get("allocation_pct", 0)):
        sym = h["symbol"]
        d = composite_results.get(sym, {})
        alloc = h.get("allocation_pct", 0)
        rs_val = h.get("allocation_rs", 0)
        total_alloc += rs_val
        ws2.append([sym, alloc / 100, rs_val, h.get("purchase_price", 0), h.get("quantity", 0), d.get("composite_score", 0), d.get("transformer_score", 0), d.get("quality_score", 0), d.get("technical_score", 0)])
        xl_r = ws2.max_row
        ws2.cell(xl_r, 1).font = _body_font(bold=True)
        ws2.cell(xl_r, 1).alignment = _left()
        ws2.cell(xl_r, 2).number_format = "0.0%"
        ws2.cell(xl_r, 3).number_format = "#,##0"
        ws2.cell(xl_r, 4).number_format = "#,##0.00"
        ws2.cell(xl_r, 5).number_format = "0"
        for sc in (6, 7, 8, 9):
            ws2.cell(xl_r, sc).number_format = "+0.0000;-0.0000;0.0000"
            ws2.cell(xl_r, sc).alignment = _centre()
        for c in range(1, len(port_hdr) + 1):
            ws2.cell(xl_r, c).border = thin_border
            ws2.cell(xl_r, c).fill = _fill(GREEN_FILL)

    port_widths = [22, 10, 14, 14, 8, 12, 14, 12, 12]
    for col_idx, width in enumerate(port_widths, 1):
        ws2.column_dimensions[get_column_letter(col_idx)].width = width
    ws2.freeze_panes = "B4"

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        wb.save(tmp.name)
    artifact = ReportArtifact.objects.create(
        kind=ReportArtifact.XLSX, portfolio=portfolio, composite_run=composite_run,
        source=ReportArtifact.Source.COMPOSITE,
    )
    now_ts = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    with open(tmp.name, 'rb') as f:
        artifact.file.save(f"{now_ts}_score_breakdown.xlsx", File(f))
    os.remove(tmp.name)
    return artifact


def save_screener_excel(results: list, min_score: int) -> "ReportArtifact | None":
    """Generates a single-sheet Excel workbook of run_screener results and saves it as a ReportArtifact.

    Companion to save_score_breakdown_excel, at a deliberately simpler scope:
    the screener is a lighter-weight PASS/FAIL/NEAR MISS filter (see
    forecasting.services.screener.get_or_screen), not the multi-layer
    composite pipeline, so this produces one styled sheet rather than the
    two-sheet Score Breakdown + Portfolio workbook. Row fill color follows
    each result's status the same way save_score_breakdown_excel colors by
    tier, for visual consistency across both reports.

    Args:
        results: List of forecasting.models.ScreenerResult instances, in the
            order they should appear on the sheet — callers (run_screener.py)
            already sort these by status/score before calling this, so no
            re-sorting happens here.
        min_score: The --min-score threshold this run was screened against,
            included in the sheet header for context (e.g. so someone
            reading the report later knows what "PASS" meant for this run).

    Returns:
        The created ReportArtifact, or None if openpyxl isn't installed or
        results is empty (nothing meaningful to write).
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        logger.warning("openpyxl not installed.")
        return None

    if not results:
        logger.warning("save_screener_excel called with no results — nothing to save.")
        return None

    import tempfile
    import os

    NAVY = "FF1E3A5F"
    WHITE = "FFFFFFFF"
    GREEN_FILL = "FFE8F5E9"
    AMBER_FILL = "FFFDF3E3"
    RED_FILL = "FFFCE4E4"
    GREY_FILL = "FFF0F0F0"
    GREEN_HDR = "FF1B5E20"
    AMBER_HDR = "FFE65100"
    RED_HDR = "FFB71C1C"
    GREY_FONT = "FF607D8B"
    STRIPE_EVEN = "FFF5F7FA"
    STRIPE_ODD = "FFFFFFFF"

    def _hdr_font(): return Font(name="Arial", bold=True, color=WHITE, size=9)
    def _body_font(bold=False, color="FF0D1B2A"): return Font(name="Arial", bold=bold, color=color, size=9)
    def _fill(argb): return PatternFill("solid", fgColor=argb)
    def _centre(): return Alignment(horizontal="center", vertical="center", wrap_text=True)
    def _left(): return Alignment(horizontal="left", vertical="center")

    thin = Side(style="thin", color="FFBDBDBD")
    thin_border = Border(left=thin, right=thin, top=thin, bottom=thin)

    status_fill = {
        "PASS": GREEN_FILL, "PASS (not selected)": GREEN_FILL,
        "NEAR MISS": AMBER_FILL, "FAIL": RED_FILL, "NO DATA": GREY_FILL,
    }
    status_font_clr = {
        "PASS": GREEN_HDR, "PASS (not selected)": GREEN_HDR,
        "NEAR MISS": AMBER_HDR, "FAIL": RED_HDR, "NO DATA": GREY_FONT,
    }

    wb = Workbook()
    ws = wb.active
    ws.title = "Screener Results"

    now_str = datetime.now().strftime("%d %b %Y  %H:%M")
    passed_count = sum(1 for r in results if r.status == "PASS")
    ws.append([
        "QuantPortfolioAI — Screener Results", f"Generated: {now_str}",
        f"Min Score: {min_score}/6", f"Symbols: {len(results)}", f"Passed: {passed_count}",
    ])
    ws.row_dimensions[1].height = 16
    ws.append([])

    HDR_ROW = 3
    headers = [
        "Symbol", "Status", "Score",
        "Trend", "Momentum", "MACD", "RSI", "Volume", "Drawdown",
        "Close", "Mom 20D", "Mom 60D", "RSI 14", "Vol Ratio",
    ]
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(HDR_ROW, c)
        cell.font = _hdr_font()
        cell.fill = _fill(NAVY)
        cell.alignment = _centre()
        cell.border = thin_border
    ws.row_dimensions[HDR_ROW].height = 24

    signal_keys = ("s1_trend", "s2_momentum", "s3_macd", "s4_rsi", "s5_volume", "s6_drawdown")

    for r in results:
        raw = r.raw_values or {}
        row_vals = [
            r.symbol.ticker, r.status, r.score,
            *["✓" if getattr(r, k) else "✗" for k in signal_keys],
            raw.get("close"), raw.get("mom_20d"), raw.get("mom_60d"),
            raw.get("rsi_14"), raw.get("vol_ratio"),
        ]
        ws.append(row_vals)
        xl_row = ws.max_row
        row_fill = _fill(status_fill.get(r.status, STRIPE_ODD))

        for col_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(xl_row, col_idx)
            cell.border = thin_border
            cell.fill = row_fill
            if col_idx == 1:
                cell.font = _body_font(bold=True)
                cell.alignment = _left()
            elif col_idx == 2:
                cell.font = Font(name="Arial", bold=True, color=status_font_clr.get(r.status, GREY_FONT), size=9)
                cell.alignment = _centre()
            elif col_idx == 3:
                cell.font = _body_font(bold=True)
                cell.alignment = _centre()
            elif col_idx <= 3 + len(signal_keys):
                cell.font = _body_font()
                cell.alignment = _centre()
            else:
                cell.alignment = _centre()
                cell.font = _body_font()
                if isinstance(val, float):
                    cell.number_format = "0.0000"

    col_widths = [16, 18, 8] + ([9] * len(signal_keys)) + [10, 10, 10, 9, 10]
    for col_idx, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A{HDR_ROW}:{get_column_letter(len(headers))}{ws.max_row}"

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        wb.save(tmp.name)
    artifact = ReportArtifact.objects.create(
        kind=ReportArtifact.XLSX, source=ReportArtifact.Source.SCREENER
    )
    now_ts = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    with open(tmp.name, 'rb') as f:
        artifact.file.save(f"{now_ts}_screener_results.xlsx", File(f))
    os.remove(tmp.name)
    return artifact